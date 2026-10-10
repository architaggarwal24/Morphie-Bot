"""DOCX loader. Word documents have no fixed page boundaries in the file
format itself (pagination happens at render time), so this always
returns a single LoadedPage - citations show the document name."""

from __future__ import annotations

import io
import logging

from .base import DocumentLoadError, DocumentLoader, LoadedPage


logger = logging.getLogger(__name__)


class DOCXLoader(DocumentLoader):
    def load(self, file_bytes: bytes) -> list[LoadedPage]:
        try:
            import docx
        except ImportError as exc:
            raise DocumentLoadError("DOCX support is not installed (python-docx).") from exc

        try:
            document = docx.Document(io.BytesIO(file_bytes))
        except Exception as exc:  # noqa: BLE001
            logger.warning("DOCX could not be parsed (%s).", type(exc).__name__)
            raise DocumentLoadError("Could not read this Word document. It may be damaged or password-protected.") from exc

        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        table_lines = []
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    table_lines.append(" | ".join(cells))

        text = "\n".join(paragraphs + table_lines).strip()
        if not text:
            raise DocumentLoadError("This Word document has no readable text.")
        return [LoadedPage(page_number=None, text=text)]
